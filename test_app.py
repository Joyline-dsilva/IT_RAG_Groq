"""No API key or external network needed. Uses a temporary database."""
import io
import json
import os
from pathlib import Path
from types import SimpleNamespace
import tempfile
import unittest
from unittest.mock import patch

import app as module
from docx import Document

class RagTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.old_db = module.DB
        module.DB = Path(self.temp.name) / 'test.db'
        with module.connect() as db:
            db.execute('CREATE TABLE documents (id INTEGER PRIMARY KEY, name TEXT, digest TEXT UNIQUE, units TEXT)')
        module.app.config['SERVER_NAME'] = 'localhost:5000'
        self.client = module.app.test_client()
        self.headers = {'X-App-Token': module.TOKEN}

    def tearDown(self):
        module.DB = self.old_db
        self.temp.cleanup()

    def post(self, route, data):
        return self.client.post(route, json=data, headers=self.headers)

    def seed(self):
        return self.post('/api/paste', {'name':'Email guide', 'text':'If Outlook is not receiving emails, check your internet connection. Restart Outlook and send a test email. If the issue continues, contact the IT support team.'})

    def test_ingest_duplicate_delete_and_unknown_question(self):
        self.assertEqual(self.seed().status_code, 200)
        self.seed()
        docs = self.client.get('/api/documents').json['documents']
        self.assertEqual(len(docs), 1)
        self.assertGreater(docs[0]['chunks'], 0)
        self.assertTrue(module.retrieve('Outlook not receiving emails'))
        self.assertFalse(module.retrieve('volcanic geology basalt'))
        self.assertEqual(self.post('/api/chat', {'question':'volcanic geology basalt'}).json['sources'], [])
        self.client.delete(f"/api/documents/{docs[0]['id']}", headers=self.headers)
        self.assertFalse(module.retrieve('Outlook'))

    def test_groq_call_has_exact_model_and_real_sources(self):
        self.seed()
        fake = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content='Restart Outlook [1].'))])
        with patch.dict(os.environ, {'GROQ_API_KEY':'test-key'}), patch.object(module, 'Groq') as mock:
            mock.return_value.chat.completions.create.return_value = fake
            response = self.post('/api/chat', {'question':'How do I fix Outlook not receiving emails?'})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json['answer'], 'Restart Outlook [1].')
            self.assertIn('Restart Outlook', response.json['sources'][0]['text'])
            call = mock.return_value.chat.completions.create.call_args.kwargs
            self.assertEqual(call['model'], 'llama-3.1-8b-instant')
            self.assertIn('REFERENCE PASSAGES', call['messages'][1]['content'])

    def test_security_and_validation(self):
        self.assertEqual(self.client.post('/api/paste', json={'text':'x'*30}).status_code, 403)
        self.assertEqual(self.client.get('/', headers={'Host':'untrusted.example'}).status_code, 403)
        self.assertEqual(self.post('/api/paste', {'text':'x'}).status_code, 400)
        self.assertEqual(self.post('/api/chat', {'question':5}).status_code, 400)
        home = self.client.get('/')
        self.assertEqual(home.status_code, 200)
        self.assertIn(b'Less searching.', home.data)
        self.assertNotIn(b'test-key', home.data)

    def test_docx_and_txt_upload(self):
        doc = Document()
        doc.add_paragraph('Printer troubleshooting procedure')
        table = doc.add_table(rows=1, cols=2)
        table.cell(0,0).text = 'Printer offline'
        table.cell(0,1).text = 'Check power and reconnect the printer cable.'
        content = io.BytesIO(); doc.save(content); content.seek(0)
        response = self.client.post('/api/upload', data={'file':(content,'printer.docx')}, headers=self.headers)
        self.assertEqual(response.status_code, 200)
        self.assertIn('reconnect', module.retrieve('Printer offline')[0]['text'])
        response = self.client.post('/api/upload', data={'file':(io.BytesIO(b'WiFi procedure: restart the wireless adapter and check the office network.'),'wifi.txt')}, headers=self.headers)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(module.retrieve('wireless adapter'))
        bad = self.client.post('/api/upload', data={'file':(io.BytesIO(b'bad'),'bad.pdf')}, headers=self.headers)
        self.assertEqual(bad.status_code, 400)

    def test_chunks_keep_end_and_overlap(self):
        text = 'network connection troubleshooting ' * 100
        parts = module.chunks(text)
        self.assertTrue(parts[-1].endswith('troubleshooting'))
        self.assertTrue(all(len(p) <= 1100 for p in parts))

if __name__ == '__main__':
    unittest.main()
